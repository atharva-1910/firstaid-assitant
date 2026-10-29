"""Builds "Raw sources and full dataset" as a Word file (export it to PDF with Word or Pages).

    backend/.venv/bin/python docs/data_article/build_data_pdf.py OUT.docx

Part A lists the raw sources (guidelines, pictures, public datasets) and the team's raw extraction
tables from docs/specification.md Section 9. Part B prints every data file the team made.
"""
import glob
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import docx
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPO = Path(__file__).resolve().parent.parent.parent
B = REPO / "backend"
OUT = sys.argv[1]
LANGS = [("en", "EN"), ("hi", "HI"), ("hinglish", "HG"), ("mr", "MR")]
NAVY = RGBColor(0x1F, 0x38, 0x64)
GREY = RGBColor(0x52, 0x51, 0x4E)


def load(p):
    return [json.loads(line) for line in open(p, encoding="utf-8") if line.strip()]


doc = docx.Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
for side in ("left_margin", "right_margin"):
    setattr(sec, side, Cm(1.8))
sec.top_margin = sec.bottom_margin = Cm(1.6)
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(9.5)
st.element.rPr.rFonts.set(qn("w:cs"), "Kohinoor Devanagari")
for name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)):
    h = doc.styles[name]
    h.font.name = "Calibri"
    h.font.size = Pt(size)
    h.font.color.rgb = NAVY

# Page numbers in the footer
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
    run = fp.add_run()
    if kind:
        el = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): kind})
    else:
        el = run._r.makeelement(qn("w:instrText"), {qn("xml:space"): "preserve"})
        el.text = text
    run._r.append(el)
    run.font.size = Pt(8)


def H(text, level=1):
    return doc.add_heading(text, level=level)


def P(text="", bold_lead=None, size=None, italic=False, color=None, after=4, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    if bold_lead:
        r = p.add_run(bold_lead + " ")
        r.bold = True
        if size:
            r.font.size = Pt(size)
    r = p.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color:
        r.font.color.rgb = color
    return p


def T(header, rows, widths_cm, size=8, header_fill="E7EBF2", chunk=100):
    """Very long tables are written as consecutive 100-row blocks: Pages will not start a table of
    hundreds of rows part-way down a page, which otherwise leaves near-empty pages."""
    rows = list(rows)
    if len(rows) > chunk:
        for k in range(0, len(rows), chunk):
            T(header, rows[k:k + chunk], widths_cm, size, header_fill, chunk)
        return None
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    tblpr = t._tbl.tblPr
    tw = tblpr.find(qn("w:tblW"))
    tw.set(qn("w:w"), str(int(sum(widths_cm) * 567)))
    tw.set(qn("w:type"), "dxa")
    for gc, wcm in zip(t._tbl.tblGrid.findall(qn("w:gridCol")), widths_cm):
        gc.set(qn("w:w"), str(int(wcm * 567)))
    hdr = t.rows[0]
    for j, h in enumerate(header):
        c = hdr.cells[j]
        c.width = Cm(widths_cm[j])
        c.paragraphs[0].add_run(str(h)).bold = True
        c.paragraphs[0].runs[0].font.size = Pt(size)
        tcpr = c._tc.get_or_add_tcPr()
        shd = tcpr.makeelement(qn("w:shd"), {qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): header_fill})
        tcpr.append(shd)
    trpr = hdr._tr.get_or_add_trPr()
    trpr.append(trpr.makeelement(qn("w:tblHeader"), {}))
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].width = Cm(widths_cm[j])
            par = cells[j].paragraphs[0]
            par.paragraph_format.space_after = Pt(0)
            run = par.add_run("" if v is None else str(v))
            run.font.size = Pt(size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def langlines(obj, size=8.5, indent=0.4):
    for code, tag in LANGS:
        v = (obj or {}).get(code)
        if v:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(indent)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(f"{tag}  ")
            r.bold = True
            r.font.size = Pt(size - 1)
            r.font.color.rgb = GREY
            r2 = p.add_run(v)
            r2.font.size = Pt(size)


# ------------------------------------------------------------------ data
recs = {Path(f).stem: json.load(open(f, encoding="utf-8")) for f in sorted(glob.glob(str(B / "content/situations/*.json")))}
ORDER = ["unresponsive", "heart_attack", "choking", "bites", "burns"]
myths = load(B / "content/myths.jsonl")
questions = json.load(open(B / "content/questions.json", encoding="utf-8"))["questions"]
strings = json.load(open(B / "content/strings.json", encoding="utf-8"))
lex = {c: json.load(open(B / f"content/lexicons/{c}.json", encoding="utf-8")) for c, _ in LANGS}
train, dev = load(B / "data/queries/train.jsonl"), load(B / "data/queries/dev.jsonl")
suite, probe = load(B / "tests/safety_suite.jsonl"), load(B / "data/queries/probe_labelled.jsonl")
ver = load(B / "data/verifier/adversarial.jsonl")
n_steps = sum(len(b["steps"]) for r in recs.values() for b in r["branches"].values())
n_dn = sum(len(b.get("do_not", [])) for r in recs.values() for b in r["branches"].values())
n_terms = sum(len(e["terms"]) for v in lex.values() for e in v["entries"])
src_count = Counter()
for r in recs.values():
    for b in r["branches"].values():
        for s in b["steps"]:
            src_count[s["source"]["doc"]] += 1

# ------------------------------------------------------------------ cover and contents
t = doc.add_paragraph()
t.alignment = WD_ALIGN_PARAGRAPH.CENTER
t.paragraph_format.space_before = Pt(120)
r = t.add_run("Multilingual First Aid Assistant")
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = NAVY
P("Raw sources and complete dataset", size=16, align=WD_ALIGN_PARAGRAPH.CENTER, after=24)
P("Taneesha Badhe · Atharva Hemade · Siddhi Karhekar · Soham Ghodake", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
P("CCNLP Mini Project · MIT World Peace University, Pune", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, after=36)
P("Repository: https://github.com/atharva-1910/firstaid-assitant", size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=GREY)
P("Every record below is printed from the repository files; nothing has been edited or sampled unless stated.", size=10,
  align=WD_ALIGN_PARAGRAPH.CENTER, italic=True, color=GREY)
page_break()

H("Contents")
T(["Part", "Section", "Source file(s)", "Records"], [
    ["A", "A1 Guideline documents (raw sources)", "—", "5 documents"],
    ["A", "A2 Raw extraction tables (from the guidelines)", "docs/specification.md, Section 9", "5 situations"],
    ["A", "A3 Other raw inputs (pictures, public datasets)", "—", "5 pictures, 2 datasets"],
    ["B", "B1 Procedure corpus (all steps, 4 languages)", "backend/content/situations/*.json", f"{n_steps} steps, {n_dn} do-not rules"],
    ["B", "B2 Harmful home remedies", "backend/content/myths.jsonl", str(len(myths))],
    ["B", "B3 Clarifying questions", "backend/content/questions.json", str(len(questions))],
    ["B", "B4 Summary and disclaimer strings", "backend/content/strings.json", "—"],
    ["B", "B5 Word lists (lexicons)", "backend/content/lexicons/*.json", f"{n_terms:,} phrases"],
    ["B", "B6 Safety suite", "backend/tests/safety_suite.jsonl", str(len(suite))],
    ["B", "B7 Probe set", "backend/data/queries/probe_labelled.jsonl", str(len(probe))],
    ["B", "B8 Template messages (train + dev)", "backend/data/queries/train.jsonl, dev.jsonl", f"{len(train) + len(dev):,}"],
    ["B", "B9 Verifier pairs", "backend/data/verifier/adversarial.jsonl", f"{len(ver):,}"],
], [1.0, 7.0, 6.4, 3.0], size=9)
P("Language tags used throughout: EN English, HI Hindi, HG Hinglish, MR Marathi. Guideline codes: G1 IFRC 2025, G2 AHA/American Red "
  "Cross 2024, G3 AHA 2025 Basic Life Support, G4 ILCOR 2025, G5 NAPSE 2024 (India).", size=9, italic=True)
page_break()

# ------------------------------------------------------------------ PART A
H("Part A. Raw sources")
P("The raw data behind the procedure corpus are five published first aid guidelines. They are copyrighted documents of their "
  "publishers and are not reproduced here; each can be downloaded free from the link given. Our steps are not copied from them word "
  "for word: each situation owner read the listed sections and rewrote every action as a short instruction, citing the document and "
  "page (A2, B1).")
H("A1 Guideline documents", 2)
T(["Code", "Document (publisher, year)", "Sections and pages used", "Our steps citing it as main source", "Link"], [
    ["G1", "International First Aid, Resuscitation and Education Guidelines 2025 (IFRC, 2026)",
     "Unresponsive and breathing p. 152; abnormal breathing / CPR p. 158; AED p. 196; choking p. 238; burns p. 288; mammal bites p. 323; insect bites and stings p. 326; snakebite p. 340; chest pain p. 358",
     str(src_count["G1"]), "https://www.ifrc.org/media/78476"],
    ["G2", "2024 American Heart Association and American Red Cross Guidelines for First Aid (Circulation 150(24), 2024)",
     "Chest pain; cooling of thermal burns; care after cooling; chemical exposures to skin and eye; snakebite; bee and wasp stings; spider and scorpion envenomation",
     str(src_count["G2"]), "https://doi.org/10.1161/CIR.0000000000001281"],
    ["G3", "2025 AHA Guidelines for CPR and ECC, Part 7: Adult Basic Life Support (AHA, 2025)",
     "Recognition of cardiac arrest; CPR; compression-only CPR; AED; foreign-body airway obstruction",
     str(src_count["G3"]), "https://cpr.heart.org/en/resuscitation-science/cpr-and-ecc-guidelines/adult-basic-life-support"],
    ["G4", "2025 International Consensus on First Aid Science with Treatment Recommendations (ILCOR, 2025)",
     "Recovery position; aspirin for chest pain; burn cooling (evidence strength only)",
     f"0 (supporting source for {sum(1 for r in recs.values() for b in r['branches'].values() for s in b['steps'] if any(a['doc'] == 'G4' for a in s.get('also_in', [])))})",
     "https://www.ilcor.org/uploads/FA-2025-COSTR-Full-Chapter.pdf"],
    ["G5", "National Action Plan for Prevention and Control of Snakebite Envenoming, NAPSE (MoHFW, Government of India, 2024)",
     "Harmful practices, p. 2", f"0 (source of {sum(1 for m in myths if m.get('source') == 'G5')} harmful-remedy entries)",
     "https://ncdc.mohfw.gov.in/wp-content/uploads/2024/07/NATIONAL-ACTION-PLAN-FOR-PREVENTION-AND-CONTROL-OF-SNAKEBITE-ENVENOMING-NAPSE.pdf"],
], [1.0, 4.4, 5.3, 2.4, 4.3], size=7.5)
if src_count.get("UNSOURCED"):
    P(f"{src_count['UNSOURCED']} steps (electrical injury) have no post-2020 source yet and are marked UNSOURCED.", size=8.5, italic=True)

H("A2 Raw extraction tables", 2)
P("These are the tables the situation owners wrote while reading the guidelines (docs/specification.md, Section 9), printed as written, "
  "with the source citation beside each step. [VERIFY] marks an item still to be checked against the PDF; (checked) marks a "
  "recommendation read in the source. The procedure corpus (B1) was converted from these tables.", size=9)
spec = (REPO / "docs/specification.md").read_text(encoding="utf-8").splitlines()
start = next(i for i, l in enumerate(spec) if l.startswith("## 9."))
end = next(i for i, l in enumerate(spec) if i > start and l.startswith("## 10"))
clean = lambda s: re.sub(r"\*\*(.+?)\*\*", r"\1", s).replace("`", "").strip()  # noqa: E731
i = start + 1
while i < end:
    line = spec[i]
    if line.startswith("|"):
        rows = []
        while i < end and spec[i].startswith("|"):
            cells = [clean(c) for c in spec[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                rows.append(cells)
            i += 1
        if rows:
            ncol = len(rows[0])
            widths = {3: [1.0, 9.6, 6.8], 2: [6.0, 11.4]}.get(ncol, [17.4 / ncol] * ncol)
            T(rows[0], [r + [""] * (ncol - len(r)) for r in rows[1:]], widths, size=7.5)
        continue
    if line.startswith("### "):
        H(clean(line[4:]), 3)
    elif line.startswith("#### "):
        P(clean(line[5:]), size=9.5, after=2).runs[0].bold = True
    elif line.strip():
        P(clean(line), size=8.5)
    i += 1

H("A3 Other raw inputs", 2)
pics = []
for r in recs.values():
    for b in r["branches"].values():
        for s in b["steps"]:
            img = s.get("image")
            if img and "wikimedia" in img.get("source_url", ""):
                pics.append([s["id"], img.get("title", ""), img["author"], img["license"], img["source_url"]])
T(["Step", "Picture (Wikimedia Commons)", "Author", "Licence", "Source"], pics, [2.0, 4.6, 3.0, 2.0, 5.8], size=7.5)
T(["Public dataset", "Use", "Link"], [
    ["XNLI (Conneau et al., 2018)", "Training data for the planned mmBERT verifier (English); not redistributed", "https://huggingface.co/datasets/facebook/xnli"],
    ["IndicXNLI (Aggarwal et al., 2022)", "Training data for the planned mmBERT verifier (Hindi, Marathi); not redistributed", "https://huggingface.co/datasets/Divyanshu/indicxnli"],
], [4.4, 7.0, 6.0], size=8)
P("The everyday phrasing in the word lists and messages comes from the four authors, one per language. The probe set was drafted "
  "with an AI assistant (B7).", size=8.5, italic=True)
page_break()

# ------------------------------------------------------------------ PART B
H("Part B. The dataset we made")
H("B1 Procedure corpus", 2)
P(f"{n_steps} steps and {n_dn} do-not rules in {sum(len(r['branches']) for r in recs.values())} branches. For each step: ID, source "
  "(document and page; 'also' lists supporting sources), the next step, and the text in four languages. Picture, timer and Yes/No "
  "check are noted where present.", size=9)
for sid in ORDER:
    r = recs[sid]
    H(f"{r['names']['en']}  ({sid}.json)", 3)
    P(f"Owner: {r.get('owner', '')} · Review status: {r.get('review_status')} · Branch on: {r.get('branch_on') or '—'}", size=8.5, color=GREY)
    for rec in r.get("recognition", []):
        P("Recognition:", size=8.5, after=0).runs[0].bold = True
        langlines(rec)
    for bn, b in r["branches"].items():
        if bn != "default":
            P(f"Branch: {bn}", size=10, after=2).runs[0].bold = True
        ur = b.get("urgency_rules")
        if ur:
            P("Urgency rules: " + "; ".join(f"if {json.dumps(u.get('if', {}), ensure_ascii=False)} → {u['urgency']}" for u in ur), size=8, color=GREY, after=2)
        for s in b["steps"]:
            src = s["source"]
            also = ", ".join(f"{a['doc']}" + (f" p. {a['page']}" if a.get("page") else "") for a in s.get("also_in", []))
            nxt = s.get("next")
            nxt = json.dumps(nxt, ensure_ascii=False) if isinstance(nxt, dict) else (nxt or "end")
            extras = [k for k in ("image", "animation", "aid", "check") if s.get(k)]
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            rr = p.add_run(f"{s['order']}. {s['id']}")
            rr.bold = True
            rr.font.size = Pt(9)
            meta = f"   source {src['doc']}" + (f" p. {src['page']}" if src.get("page") else "") + (f" · also {also}" if also else "") + f" · next {nxt}"
            if extras:
                meta += " · has " + ", ".join(extras)
            m = p.add_run(meta)
            m.font.size = Pt(7.5)
            m.font.color.rgb = GREY
            langlines(s["text"])
        if b.get("do_not"):
            P("Do not:", size=9, after=0).runs[0].bold = True
            for d in b["do_not"]:
                langlines(d["text"])
                if d.get("reason") and d["reason"].get("en"):
                    P("Reason:", size=8, after=0, color=GREY)
                    langlines(d["reason"], size=8, indent=0.8)
    page_break()

H("B2 Harmful home remedies", 2)
for m in myths:
    p = P(f"{m['id']}  ·  {m['situation']} / {', '.join(m.get('subtypes', []))}  ·  source {m.get('source')}", size=9, after=0)
    p.runs[0].bold = True
    langlines(m["action"])
    P("Why harmful:", size=8, after=0, color=GREY)
    langlines(m["reason"], size=8, indent=0.8)
H("B3 Clarifying questions", 2)
for q in questions:
    P(f"{q['id']}" + (f"  (safety fact: {q['safety_slot']})" if q.get("safety_slot") else ""), size=9, after=0).runs[0].bold = True
    langlines(q["text"])
    T(["Option", "EN", "HI", "HG", "MR", "Sets"], [[o["id"]] + [o["label"].get(c, "") for c, _ in LANGS] + [json.dumps(o.get("set", {}), ensure_ascii=False)]
                                                  for o in q["options"]], [2.0, 3.0, 3.0, 3.0, 3.0, 3.4], size=7.5)
H("B4 Summary and disclaimer strings", 2).paragraph_format.keep_with_next = False
rows = []


def flat(prefix, o):
    if isinstance(o, dict) and "en" in o:
        rows.append([prefix] + [o.get(c, "") for c, _ in LANGS])
    elif isinstance(o, dict):
        for k, v in o.items():
            if not k.startswith("_"):
                flat(f"{prefix}.{k}" if prefix else k, v)


flat("", strings)
T(["Key", "EN", "HI", "HG", "MR"], rows, [3.0, 3.6, 3.6, 3.6, 3.6], size=7)
page_break()

H("B5 Word lists (lexicons)", 2)
P("For each language: every entry (the phrases, and the facts they set when affirmed or negated), then the negation and uncertainty cues "
  "and the action phrases. A trailing * matches any ending.", size=9)
for code, tag in LANGS:
    L = lex[code]
    H(f"{tag} — {code}.json ({len(L['entries'])} entries, {sum(len(e['terms']) for e in L['entries'])} phrases; cue direction: {L.get('cue_direction')})", 3)
    P("Negation cues: " + ", ".join(L.get("negation_cues", [])), size=8)
    P("Uncertainty cues: " + ", ".join(L.get("uncertainty_cues", [])), size=8)
    if L.get("scope_breakers"):
        P("Scope breakers: " + ", ".join(L["scope_breakers"]), size=8)
    rows = []
    for e in L["entries"]:
        sets = json.dumps(e.get("set", {}), ensure_ascii=False)
        extra = []
        if e.get("neg"):
            extra.append("neg→" + json.dumps(e["neg"], ensure_ascii=False))
        if e.get("flags"):
            extra.append("flags " + ",".join(e["flags"]))
        if e.get("sit"):
            extra.append("sit " + ",".join(e["sit"]))
        rows.append([", ".join(e["terms"]), sets + ("  " + "; ".join(extra) if extra else "")])
    T(["Phrases", "Sets (and on negation)"], rows, [11.0, 6.4], size=7)
    acts = L.get("actions", [])
    if acts:
        T(["Action phrases", "Action"], [[", ".join(a["terms"]), a["action"]] for a in acts], [12.4, 5.0], size=7)
page_break()

H("B6 Safety suite", 2)
P(f"{len(suite)} team-written messages; {sum(c['must_be_emergency'] for c in suite)} must be treated as emergencies.", size=9)
T(["ID", "Lang", "Message", "Must be emergency", "Expected"], [[c["id"], c["lang"], c["text"], "yes" if c["must_be_emergency"] else "no", c["expect"]] for c in suite],
  [3.4, 1.4, 8.6, 1.8, 2.2], size=7.5)
H("B7 Probe set", 2)
P(f"{len(probe)} short, misspelt, code-mixed messages drafted with an AI assistant; situation and urgency are silver labels to be checked by the team.", size=9)
T(["ID", "Lang", "Message", "Situation", "Subtype", "Urgency"], [[r["id"], r["lang"], r["text"], r.get("situation") or "—", r.get("subtype") or "—", r.get("urgency") or "—"] for r in probe],
  [3.2, 1.4, 6.8, 2.2, 2.2, 1.6], size=7.5)
page_break()

H("B8 Template messages", 2).paragraph_format.keep_with_next = False
P(f"{len(train):,} training and {len(dev)} development messages, 150 per class per language, generated by backend/scripts/generate_queries.py. "
  "Slots are the facts the message states.", size=9)
fmt_slots = lambda s: "; ".join(f"{k}={v}" for k, v in (s or {}).items())  # noqa: E731
for split_name, rows_ in (("train.jsonl", train), ("dev.jsonl", dev)):
    for code, tag in LANGS:
        sub = [r for r in rows_ if r["lang"] == code]
        h = H(f"{split_name} · {tag} ({len(sub)} messages)", 3)
        h.paragraph_format.keep_with_next = False  # a 720-row table cannot stay with its heading; let it flow
        T(["ID", "Message", "Situation / subtype", "Urgency", "Slots"],
          [[r["id"], r["text"], r["situation"] + (f" / {r['subtype']}" if r.get("subtype") else ""), r.get("urgency") or "—", fmt_slots(r.get("slots"))] for r in sub],
          [2.4, 6.6, 3.0, 1.6, 3.8], size=6.5)
page_break()

H("B9 Verifier pairs", 2)
P(f"{len(ver):,} pairs, grouped by the step they were made from. 'same' pairs keep the meaning; 'different' pairs change it.", size=9)
groups = defaultdict(list)
for v in ver:
    groups[(v["source"], v["premise"])].append(v)
for (src, prem), items in groups.items():
    p = P(f"{src}: ", size=8, after=0)
    p.runs[0].bold = True
    p.add_run(prem).font.size = Pt(8)
    T(["Kind", "Label", "Candidate"], [[v["kind"], v["label"], v["hypothesis"]] for v in items], [2.8, 1.6, 13.0], size=7)

zoom = doc.settings.element.find(qn("w:zoom"))
if zoom is not None and zoom.get(qn("w:percent")) is None:
    zoom.set(qn("w:percent"), "100")
doc.core_properties.title = "Multilingual First Aid Assistant — raw sources and complete dataset"
doc.save(OUT)
print("wrote", OUT, "| steps", n_steps, "| messages", len(train) + len(dev), "| pairs", len(ver), "| verifier groups", len(groups))
